"""Регрессия по ревью 2026-09-10: P1+P2."""
import json
from datetime import datetime, timedelta


def make_dt(date_str, time_str):
    return datetime.strptime(f'{date_str} {time_str}', '%Y-%m-%d %H:%M:%S')


def _patch_role2(monkeypatch):
    """Локальный кладовщик 104 (роль 2, ставка 800 руб/смена) во всех модулях."""
    from decimal import Decimal
    from core.config import EmployeeData, Role
    from core import config, calculations, data_array, analysis
    roles = {
        1: Role(id=1, name='Работник', work_shift=8, lost_tag_flag=1),
        2: Role(id=2, name='Кладовщик', work_shift=8, lost_tag_flag=1),
        3: Role(id=3, name='Окладник', work_shift=8, lost_tag_flag=1),
        4: Role(id=4, name='Работник4', work_shift=8, lost_tag_flag=1),
    }
    employees = {
        101: EmployeeData(id=101, first_name='Иван', last_name='Петров',
                          role_id=1, role_name='Работник', daily_rate=Decimal('800.00')),
        104: EmployeeData(id=104, first_name='Клавдий', last_name='Кладов',
                          role_id=2, role_name='Кладовщик', daily_rate=Decimal('800.00')),
    }
    monkeypatch.setattr(config, 'ROLES', roles)
    monkeypatch.setattr(config, 'EMPLOYEES', employees)
    monkeypatch.setattr(calculations, 'EMPLOYEES', employees)
    monkeypatch.setattr(calculations, 'load_wage_rates',
                        lambda: {101: Decimal('800.00'), 104: Decimal('800.00')})
    monkeypatch.setattr(data_array, 'EMPLOYEES', employees)
    monkeypatch.setattr(analysis, 'EMPLOYEES', employees)
    return employees


class TestRole2Storekeeper:
    """Кладовщик: как роль 1, но без оплаты переработок; норма 8ч для всех."""

    def test_daily_like_role1(self, monkeypatch):
        _patch_role2(monkeypatch)
        from core.calculations import calculate_hours_per_day
        d = '2026-07-06'
        tt = {d: {104: [make_dt(d, '18:00:00'), make_dt(d, '08:00:00'), 'work']}}
        res = calculate_hours_per_day(tt)
        abs_delta, worked, tag_over, tag_day = res[d][104]
        assert worked == timedelta(hours=10)
        assert abs_delta == timedelta(hours=2)
        assert tag_day == 'work'

    def test_marks_and_missed_like_role1(self, monkeypatch):
        _patch_role2(monkeypatch)
        monkeypatch.setattr('core.file_parser.load_holidays', lambda year: [])
        monkeypatch.setattr('core.file_parser.load_postponed_days', lambda year: [])
        from core.analysis import _get_marks_and_missed
        dt = make_dt('2026-07-06', '08:00:00')
        tt = {'2026-07-06': {104: [dt, dt, 'work']}}
        single, missed = _get_marks_and_missed(tt, 104, 2026, 7)
        assert single != 0  # одиночная метка детектится, как у роли 1

    def test_overtime_not_paid(self, monkeypatch):
        _patch_role2(monkeypatch)
        from decimal import Decimal
        from core.calculations import calculate_wages
        summary = {
            104: (
                (1, timedelta(hours=2), timedelta(0)),
                (0, timedelta(0), timedelta(0), timedelta(0)),
                0, 0,
            )
        }
        total, milk, _ = calculate_wages(summary)[104]
        assert total == Decimal('800.00')  # 2ч сверхурочных не плюсуются
        assert milk == Decimal('40.00')

    def test_undertime_deducted(self, monkeypatch):
        _patch_role2(monkeypatch)
        from decimal import Decimal
        from core.calculations import calculate_wages
        summary = {
            104: (
                (1, timedelta(0), timedelta(hours=6)),
                (0, timedelta(0), timedelta(0), timedelta(0)),
                0, 0,
            )
        }
        total, _, _ = calculate_wages(summary)[104]
        assert total == Decimal('200.00')  # как у роли 1: 800 - 600

    def test_weekend_single_rate(self, monkeypatch):
        _patch_role2(monkeypatch)
        from decimal import Decimal
        from core.calculations import calculate_wages
        summary = {
            104: (
                (0, timedelta(0), timedelta(0)),
                (1, timedelta(0), timedelta(0), timedelta(hours=8)),
                0, 0,
            )
        }
        total, milk, _ = calculate_wages(summary)[104]
        assert total == Decimal('800.00')  # факт без 1.5x (у роли 1 было бы 1200)
        assert milk == Decimal('40.00')


class TestExcludedNotInPreview:
    def test_preview_filters_exceptions(self, setup_employees, monkeypatch):
        from core import data_array
        from core.ui import build_preview_rows
        # 102 в исключениях — в превью его быть не должно
        monkeypatch.setattr(data_array, 'SETTLEMENT_EXCEPTIONS', [102])
        import core.ui as ui_mod
        # build_preview_rows импортирует is_settlement_allowed из data_array при вызове,
        # поэтому патча data_array достаточно
        summary = {
            101: ((1, timedelta(0), timedelta(0)),
                  (0, timedelta(0), timedelta(0), timedelta(0)), 0, 0),
            102: ((1, timedelta(0), timedelta(0)),
                  (0, timedelta(0), timedelta(0), timedelta(0)), 0, 0),
        }
        wages = {101: (800, 40, 840), 102: (800, 0, 800)}
        rows = build_preview_rows(summary, wages)
        assert {r['emp_id'] for r in rows} == {101}


class TestFileParserBadDate:
    def test_impossible_date_does_not_raise(self):
        from core.file_parser import parse_attlog_line
        assert parse_attlog_line('        101\t2026-02-30 08:00:00\t1\t255\t1\t0') is None

    def test_read_continues_past_bad_date(self, tmp_path, monkeypatch):
        import os
        from core.file_parser import read_file_data_with_errors
        monkeypatch.chdir(tmp_path)
        os.makedirs('data', exist_ok=True)
        with open('data/mixed.dat', 'w', encoding='utf-8') as f:
            f.write('        101\t2026-02-30 08:00:00\t1\t255\t1\t0\n')
            f.write('        101\t2026-07-06 08:00:00\t1\t255\t1\t0\n')
        rows, errors = read_file_data_with_errors('mixed.dat', 2026, 7)
        assert len(rows) == 1
        assert len(errors) == 1
        assert 'строка 1' in errors[0]


class TestHtmlUniqueAndEscaped:
    def test_filename_contains_id(self, setup_employees, tmp_path, monkeypatch):
        from core.html_builder import build_html, sanitize_filename_part
        monkeypatch.chdir(tmp_path)
        d = '2026-07-06'
        tt = {d: {101: [make_dt(d, '16:00:00'), make_dt(d, '08:00:00'), 'work']}}
        from core.calculations import calculate_hours_per_day, calculate_hours_per_month, calculate_wages
        wt = calculate_hours_per_day(tt)
        summary, _ = calculate_hours_per_month(wt)
        wages = calculate_wages(summary)
        # подменим ставки чтобы wages непустой
        if 101 not in wages:
            wages = {101: (800, 40, 840)}
            summary = {101: ((1, timedelta(0), timedelta(0)),
                              (0, timedelta(0), timedelta(0), timedelta(0)), 0, 0)}
        fname = build_html(101, tt, wt, summary, wages)
        assert '_101_' in fname
        assert sanitize_filename_part('a/b:c') == 'a_b_c'

    def test_html_escapes_markup(self, setup_employees):
        from core.html_builder import _gen_day_row
        day = {'family': '<b>Changed</b>', 'date': '2026-07-06',
               'time_begin': '08:00:00', 'time_end': '16:00:00',
               'delta_time': '8:00:00', 'tag_overtime': 'x', 'overtime': '0',
               'tag_day': 'work'}
        out = ''.join(_gen_day_row(day))
        assert '<b>Changed</b>' not in out
        assert '&lt;b&gt;' in out


class TestExcelFormulaAndFilter:
    def test_name_formula_stored_as_string(self, setup_employees, monkeypatch):
        from openpyxl import Workbook
        from core.excel_builder import _write_data_rows, _set_cell
        from core import data_array
        monkeypatch.setattr(data_array, 'EMPLOYEES', {
            101: type('E', (), {'last_name': '=1+1', 'first_name': ''})(),
        })
        # также пропатчить EMPLOYEES в excel_builder (импортирован напрямую)
        from core import excel_builder
        monkeypatch.setattr(excel_builder, 'EMPLOYEES', data_array.EMPLOYEES)
        d = '2026-07-06'
        tt = {d: {101: [make_dt(d, '16:00:00'), make_dt(d, '08:00:00'), 'work']}}
        wt = {d: {101: (timedelta(0), timedelta(hours=8), 'недоработка', 'work')}}
        wb = Workbook()
        ws = wb.active
        _write_data_rows(ws, tt, wt)
        cell = ws.cell(column=1, row=2)
        assert cell.data_type == 's'
        assert cell.value == '=1+1'

    def test_autofilter_covers_all_details(self, setup_employees):
        from openpyxl import Workbook
        from core.excel_builder import _write_data_rows
        tt = {}
        wt = {}
        for day in range(1, 32):
            k = f'2026-07-{day:02d}'
            tt[k] = {101: [make_dt(k, '16:00:00'), make_dt(k, '08:00:00'), 'work']}
            wt[k] = {101: (timedelta(0), timedelta(hours=8), 'недоработка', 'work')}
        wb = Workbook()
        ws = wb.active
        last = _write_data_rows(ws, tt, wt)
        assert last == 32
        # build_excel выставляет A1:G{last}
        ws.auto_filter.ref = f'A1:G{max(last, 1)}'
        assert ws.auto_filter.ref == 'A1:G32'

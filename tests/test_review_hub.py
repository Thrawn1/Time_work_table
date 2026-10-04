"""Обзор месяца: средние времена, автозаполнение, больничный, ремонтники, ручное исключение.

Режим работы оператора: ремонтники — строкой «Имя Фамилия — дни месяца»; у цеховых —
число проблем и среднее время прихода/ухода; сотрудника можно исключить из расчёта;
проблемы закрываются по одной или автоматически (среднее ± 5 минут).
"""
import random
from datetime import datetime, timedelta
from decimal import Decimal
from unittest.mock import patch

import pytest

from core import analysis, config
from core.analysis import (
    AverageTimes,
    auto_complete_day,
    auto_complete_single,
    auto_fill_singles,
    compute_average_times,
    format_average,
)
from core.day_models import DayMark


def make_dt(date_str, time_str):
    return datetime.strptime(f'{date_str} {time_str}', '%Y-%m-%d %H:%M:%S')


def sec(h, m=0, s=0):
    return h * 3600 + m * 60 + s


def day(date_key, come, go, tag='work'):
    return DayMark(go=make_dt(date_key, go), come=make_dt(date_key, come), tag=tag)


@pytest.fixture(autouse=True)
def _clean_state():
    """Журнал и ручные исключения — глобальное состояние модулей: чистим вокруг теста."""
    analysis.clear_journal()
    config.MANUAL_EXCLUSIONS.clear()
    yield
    analysis.clear_journal()
    config.MANUAL_EXCLUSIONS.clear()


def feed(*answers):
    """input(): ответы по очереди, затем EOFError (закрытый ввод), а не StopIteration."""
    it = iter(answers)

    def _input(prompt=''):
        try:
            return next(it)
        except StopIteration:
            raise EOFError from None

    return _input


class TestAverages:
    def test_average_over_complete_workdays(self):
        tt = {
            '2026-07-06': {101: day('2026-07-06', '08:00:00', '17:00:00')},
            '2026-07-07': {101: day('2026-07-07', '08:10:00', '17:10:00')},
        }
        avg = compute_average_times(tt, 101)
        assert avg == AverageTimes(come=sec(8, 5), go=sec(17, 5), days=2)
        assert avg.come_text == '08:05:00' and avg.go_text == '17:05:00'

    def test_ignores_double_punch_weekend_single_and_skipped_dates(self):
        tt = {
            '2026-07-06': {101: day('2026-07-06', '08:00:00', '17:00:00')},
            # двойное касание: «смена» в 7 секунд не должна тянуть среднее ухода
            '2026-07-07': {101: day('2026-07-07', '08:00:00', '08:00:07')},
            '2026-07-08': {101: day('2026-07-08', '08:00:00', '08:00:00')},
            '2026-07-11': {101: day('2026-07-11', '11:00:00', '15:00:00', tag='weekend')},
            '2026-07-09': {101: day('2026-07-09', '09:00:00', '18:00:00')},
        }
        assert compute_average_times(tt, 101).days == 2
        # день, уже правленный оператором, в среднее не входит
        assert compute_average_times(tt, 101, skip_dates={'2026-07-09'}).days == 1

    def test_none_without_complete_days(self):
        tt = {'2026-07-08': {101: day('2026-07-08', '08:00:00', '08:00:00')}}
        assert compute_average_times(tt, 101) is None
        assert compute_average_times({}, 101) is None
        assert 'нет данных' in format_average(None)

    def test_own_edits_do_not_shift_average(self, setup_employees):
        tt = {
            '2026-07-06': {101: day('2026-07-06', '08:00:00', '17:00:00')},
            '2026-07-07': {101: day('2026-07-07', '10:00:00', '19:00:00')},
        }
        analysis.record_edit(101, '2026-07-07', 'рабочие дни x1', None,
                             tt['2026-07-07'][101], employees=setup_employees)
        assert analysis._current_average(tt, 101).come == sec(8)


class TestAutoFill:
    def test_random_time_stays_within_five_minutes_and_day(self):
        from randomazer_time_value import random_time_near

        rng = random.Random(7)
        for _ in range(300):
            assert abs(random_time_near(sec(8, 30), rng=rng) - sec(8, 30)) <= 300
        for _ in range(100):
            assert 0 <= random_time_near(60, rng=rng) <= 360
            assert sec(23, 55) <= random_time_near(86399, rng=rng) <= 86399

    def test_single_arrival_gets_departure_near_average(self):
        avg = AverageTimes(come=sec(8), go=sec(17), days=10)
        rng = random.Random(1)
        for _ in range(50):
            come, go, side = auto_complete_single(make_dt('2026-07-06', '08:03:00'), avg, rng=rng)
            assert side == 'уход' and come == make_dt('2026-07-06', '08:03:00')
            assert abs((go - make_dt('2026-07-06', '17:00:00')).total_seconds()) <= 300

    def test_single_departure_gets_arrival_near_average(self):
        avg = AverageTimes(come=sec(8), go=sec(17), days=10)
        rng = random.Random(2)
        for _ in range(50):
            come, go, side = auto_complete_single(make_dt('2026-07-06', '17:20:00'), avg, rng=rng)
            assert side == 'приход' and go == make_dt('2026-07-06', '17:20:00')
            assert abs((come - make_dt('2026-07-06', '08:00:00')).total_seconds()) <= 300

    def test_whole_day_from_averages(self):
        avg = AverageTimes(come=sec(8), go=sec(17), days=10)
        come, go = auto_complete_day('2026-07-21', avg, rng=random.Random(3))
        assert come.date() == go.date() == datetime(2026, 7, 21).date()
        assert abs((come - make_dt('2026-07-21', '08:00:00')).total_seconds()) <= 300
        assert abs((go - make_dt('2026-07-21', '17:00:00')).total_seconds()) <= 300


def _full_days(*dates, come='08:00:00', go='17:00:00'):
    return {d: {101: day(d, come, go)} for d in dates}


class TestEditMenus:
    def test_single_mark_option_3_autofills_and_journals(self, setup_employees):
        tt = _full_days('2026-07-06', '2026-07-07', '2026-07-08')
        tt['2026-07-09'] = {101: day('2026-07-09', '08:02:00', '08:02:00')}
        single = [['2026-07-09', make_dt('2026-07-09', '08:02:00')]]
        with patch.object(analysis, '_get_marks_and_missed', return_value=(single, 0)):
            with patch('builtins.input', feed('3', 'д')):
                with patch.object(analysis, '_save_session', lambda x: None):
                    analysis.analyze_for_edit(tt, 101, 2026, 7)
        mark = tt['2026-07-09'][101]
        assert mark.come == make_dt('2026-07-09', '08:02:00')
        assert abs((mark.go - make_dt('2026-07-09', '17:00:00')).total_seconds()) <= 300
        entries = analysis.get_journal()
        assert len(entries) == 1
        assert entries[0]['action'] == 'одиночная метка (авто по среднему)'
        assert entries[0]['randomized'] is True

    def test_single_mark_autofill_declined_changes_nothing(self, setup_employees):
        tt = _full_days('2026-07-06', '2026-07-07')
        tt['2026-07-09'] = {101: day('2026-07-09', '08:02:00', '08:02:00')}
        single = [['2026-07-09', make_dt('2026-07-09', '08:02:00')]]
        with patch.object(analysis, '_get_marks_and_missed', return_value=(single, 0)):
            with patch('builtins.input', feed('3', 'н', '0')):
                with patch.object(analysis, '_save_session', lambda x: None):
                    analysis.analyze_for_edit(tt, 101, 2026, 7)
        assert tt['2026-07-09'][101].go == tt['2026-07-09'][101].come
        assert analysis.get_journal() == []

    def test_missed_days_option_6_autofills_every_day(self, setup_employees):
        tt = _full_days('2026-07-06', '2026-07-07')
        missed = ['2026-07-21', '2026-07-22']
        with patch.object(analysis, '_get_marks_and_missed', return_value=(0, missed)):
            with patch('builtins.input', feed('6', 'д')):
                with patch.object(analysis, '_save_session', lambda x: None):
                    analysis.analyze_for_edit(tt, 101, 2026, 7)
        for d in missed:
            mark = tt[d][101]
            assert mark.tag == 'work'
            assert abs((mark.come - make_dt(d, '08:00:00')).total_seconds()) <= 300
            assert abs((mark.go - make_dt(d, '17:00:00')).total_seconds()) <= 300
        entries = analysis.get_journal()
        assert [e['date'] for e in entries] == missed
        assert all(e['action'] == 'рабочие дни x2 (авто по среднему)' and e['randomized']
                   for e in entries)

    def test_autofill_unavailable_without_average(self, setup_employees):
        tt: dict = {}
        with patch.object(analysis, '_get_marks_and_missed', return_value=(0, ['2026-07-21'])):
            with patch('builtins.input', feed('6', '0')):
                with patch.object(analysis, '_save_session', lambda x: None):
                    analysis.analyze_for_edit(tt, 101, 2026, 7)
        assert tt == {}
        assert analysis.get_journal() == []

    def test_sick_leave_for_range(self, setup_employees):
        tt: dict = {}
        missed = ['2026-07-21', '2026-07-22']
        with patch.object(analysis, '_get_marks_and_missed', return_value=(0, missed)):
            with patch('builtins.input', feed('5', 'д')):
                with patch.object(analysis, '_save_session', lambda x: None):
                    analysis.analyze_for_edit(tt, 101, 2026, 7)
        for d in missed:
            assert tt[d][101].tag == 'sick'
        assert [e['action'] for e in analysis.get_journal()] == ['больничный x2'] * 2

    def test_per_day_menu_has_sick_and_auto(self, setup_employees):
        tt = _full_days('2026-07-06', '2026-07-07')
        missed = ['2026-07-21', '2026-07-22']
        # [4] разобрать по одному: первый день — больничный, второй — автозаполнение
        with patch.object(analysis, '_get_marks_and_missed', return_value=(0, missed)):
            with patch('builtins.input', feed('4', '5', 'д', '6', 'д')):
                with patch.object(analysis, '_save_session', lambda x: None):
                    analysis.analyze_for_edit(tt, 101, 2026, 7)
        assert tt['2026-07-21'][101].tag == 'sick'
        assert tt['2026-07-22'][101].tag == 'work'

    def test_bulk_autofill_all_singles_one_confirmation_one_save(self, setup_employees):
        tt = _full_days('2026-07-06', '2026-07-07', '2026-07-08')
        tt['2026-07-09'] = {101: day('2026-07-09', '08:02:00', '08:02:00')}
        tt['2026-07-10'] = {101: day('2026-07-10', '17:03:00', '17:03:00')}
        singles = [['2026-07-09', make_dt('2026-07-09', '08:02:00')],
                   ['2026-07-10', make_dt('2026-07-10', '17:03:00')]]
        saves = []
        with patch.object(analysis, '_get_marks_and_missed', return_value=(singles, 0)):
            with patch('builtins.input', feed('д')):
                with patch.object(analysis, '_save_session', lambda x: saves.append(1)):
                    done = auto_fill_singles(tt, 101, 2026, 7)
        assert done == 2 and len(saves) == 1
        assert tt['2026-07-09'][101].go > tt['2026-07-09'][101].come
        assert tt['2026-07-10'][101].go > tt['2026-07-10'][101].come
        assert len(analysis.get_journal()) == 2

    def test_bulk_autofill_without_average_does_nothing(self, setup_employees):
        tt = {'2026-07-09': {101: day('2026-07-09', '08:02:00', '08:02:00')}}
        singles = [['2026-07-09', make_dt('2026-07-09', '08:02:00')]]
        with patch.object(analysis, '_get_marks_and_missed', return_value=(singles, 0)):
            assert auto_fill_singles(tt, 101, 2026, 7) == 0
        assert analysis.get_journal() == []


class TestSickLeaveCalc:
    def test_hours_and_summary_count_sick_separately(self, setup_employees):
        from core.calculations import calculate_hours_per_day, calculate_hours_per_month

        tt = {
            '2026-07-06': {101: day('2026-07-06', '08:00:00', '16:00:00')},
            '2026-07-07': {101: day('2026-07-07', '00:00:02', '00:00:02', tag='sick')},
        }
        work_time = calculate_hours_per_day(tt)
        assert work_time['2026-07-07'][101].day_tag == 'sick'
        summary, restructured = calculate_hours_per_month(work_time)
        month = summary[101]
        assert month.sick_days == 1 and month.vacation_days == 0 and month.truancy_days == 0
        assert month.work.days == 1
        assert [len(g) for g in restructured[101]] == [1, 0, 0, 0]
        # legacy-протокол индексов не изменился
        assert len(month) == 4 and month[2] == 0 and month[3] == 0

    def test_legacy_wages_do_not_pay_sick_days(self, setup_employees):
        from core.calculations import calculate_hours_per_day, calculate_hours_per_month, calculate_wages

        base = {'2026-07-06': {101: day('2026-07-06', '08:00:00', '16:00:00')}}
        with_sick = dict(base)
        with_sick['2026-07-07'] = {101: day('2026-07-07', '00:00:02', '00:00:02', tag='sick')}
        rates = {101: Decimal('800.00')}

        def total(tt):
            summary, _ = calculate_hours_per_month(calculate_hours_per_day(tt))
            return calculate_wages(summary, rates=rates, employees=setup_employees)[101]

        assert total(with_sick).salary == total(base).salary == Decimal('800.00')

    def test_new_model_blocks_full_month_bonus_with_reason(self):
        from core.pay_calc import PayInputs, calculate_pay

        inputs = PayInputs(
            monthly_base=Decimal('60000'), base_day_hours=8, full_month_bonus=Decimal('5000'),
            workdays=20, shift_norm_hours=Decimal('8'), fact_hours=Decimal('152'),
            workdays_present=19, sick_days=1, overtime_eligible=True,
            full_month_eligible=True)
        result = calculate_pay(inputs)
        assert result.full_month_ok is False
        assert result.full_month_reason == 'больничный: 1 дн.'
        assert result.ordinary_pay == Decimal('57000.00')  # дни больничного не оплачены

    def test_session_accepts_sick_tag(self):
        from core.session import validate_session_raw

        raw = {'version': 2, 'entries': {'2026-07-07': {
            '101': ['2026-07-07T00:00:02', '2026-07-07T00:00:02', 'sick']}}}
        table, errors = validate_session_raw(raw)
        assert errors == [] and table['2026-07-07'][101].tag == 'sick'

    def test_reports_label_sick_day(self, setup_employees):
        from core.calculations import calculate_hours_per_day
        from core.html_builder import _build_daily_data

        tt = {'2026-07-07': {101: day('2026-07-07', '00:00:02', '00:00:02', tag='sick')}}
        rows = _build_daily_data(101, tt, calculate_hours_per_day(tt), setup_employees)
        assert rows[0]['tag_day'] == 'Больничный'


class TestRepairmen:
    def _staff(self):
        E = config.EmployeeData
        return {
            1: E(1, 'Иван', 'Иванов', 1, 'Работник'),
            4: E(4, 'Максим', 'Смирнов', 4, 'Ремонтники'),
            5: E(5, 'Анна', 'Белова', 4, 'Ремонтники'),
        }

    def test_line_per_repairman_with_sorted_days_and_genitive_month(self):
        from core.ui import build_repairmen_lines

        tt = {f'2026-09-{d:02d}': {4: day(f'2026-09-{d:02d}', '09:00:00', '09:00:00')}
              for d in (25, 12, 18, 15)}
        tt['2026-09-01'] = {1: day('2026-09-01', '08:00:00', '17:00:00')}
        lines = build_lines = build_repairmen_lines(tt, self._staff(), 2026, 9)
        assert lines == ['Анна Белова — выходов нет',
                         'Максим Смирнов — 12, 15, 18, 25 сентября']
        assert build_lines is lines

    def test_no_repairmen_no_lines_and_nothing_printed(self, capsys):
        from core.ui import build_repairmen_lines, print_repairmen

        staff = {1: self._staff()[1]}
        assert build_repairmen_lines({}, staff, 2026, 9) == []
        print_repairmen([])
        assert capsys.readouterr().out == ''

    def test_print_shows_lines(self, capsys):
        from core.ui import print_repairmen

        print_repairmen(['Максим Смирнов — 12, 15 сентября'])
        assert 'Максим Смирнов — 12, 15 сентября' in capsys.readouterr().out


class TestManualExclusion:
    def test_toggle_excludes_everywhere_and_returns(self, setup_employees):
        from core.data_array import exclusion_reason, is_included_in_settlement, is_settlement_allowed
        from core.review import toggle_manual_exclusion

        assert is_included_in_settlement(101, setup_employees)
        assert toggle_manual_exclusion(101) is True
        assert exclusion_reason(101, setup_employees) == config.MANUAL_EXCLUSION_REASON
        assert exclusion_reason(101, setup_employees, include_manual=False) == ''
        assert not is_included_in_settlement(101, setup_employees)
        assert not is_settlement_allowed(101)
        assert toggle_manual_exclusion(101) is False
        assert is_included_in_settlement(101, setup_employees)

    def test_preserve_config_restores_manual_exclusions(self):
        with config.preserve_config():
            config.MANUAL_EXCLUSIONS[7] = 'x'
        assert config.MANUAL_EXCLUSIONS == {}

    def test_dashboard_marks_manual_and_keeps_numbers_stable(
            self, setup_employees, mock_holidays_jan2026, mock_postponed_empty):
        from core.ui import build_dashboard_rows, review_numbering

        tt = {'2026-07-06': {101: day('2026-07-06', '08:00:00', '17:00:00'),
                             102: day('2026-07-06', '08:00:00', '17:00:00')}}
        config.MANUAL_EXCLUSIONS[101] = config.MANUAL_EXCLUSION_REASON
        rows = build_dashboard_rows(tt, [101, 102, 103], 2026, 7, employees=setup_employees)
        by_id = {r['emp_id']: r for r in rows}
        assert by_id[101]['manual'] and by_id[101]['excluded']
        assert not by_id[102]['manual'] and not by_id[102]['excluded']
        assert by_id[102]['avg_come'] == '08:00:00' and by_id[102]['avg_go'] == '17:00:00'
        # 103 исключён правилами роли — по номеру недоступен, ручным не считается
        assert not by_id[103]['manual']
        # номера — в порядке справочника и не зависят от исключения
        assert [r['emp_id'] for r in review_numbering(rows)] == [101, 102]
        config.MANUAL_EXCLUSIONS.clear()
        rows = build_dashboard_rows(tt, [101, 102, 103], 2026, 7, employees=setup_employees)
        assert [r['emp_id'] for r in review_numbering(rows)] == [101, 102]


class TestRunReview:
    def _table(self):
        return {
            '2026-07-06': {101: day('2026-07-06', '08:00:00', '17:00:00'),
                           102: day('2026-07-06', '08:00:00', '17:00:00')},
            '2026-07-07': {101: day('2026-07-07', '08:10:00', '17:10:00')},
        }

    def test_exclude_and_return_call_callback(self, setup_employees, capsys,
                                              mock_holidays_jan2026, mock_postponed_empty):
        from core.review import current_settlement_ids, run_review

        calls = []
        # x1 (латинская) исключает №1, х1 (русская «х») возвращает его же: номера стабильны
        with patch('builtins.input', feed('x1', 'х1', '')):
            run_review(self._table(), setup_employees, [101, 102], 2026, 7,
                       on_exclusion_change=lambda: calls.append(dict(config.MANUAL_EXCLUSIONS)))
        assert calls == [{101: config.MANUAL_EXCLUSION_REASON}, {}]
        assert current_settlement_ids([101, 102], setup_employees) == [101, 102]
        assert 'исключён из расчёта' in capsys.readouterr().out

    def test_excluded_stay_excluded_from_settlement_ids(self, setup_employees,
                                                        mock_holidays_jan2026, mock_postponed_empty):
        from core.review import current_settlement_ids, run_review

        with patch('builtins.input', feed('x2', '')):
            run_review(self._table(), setup_employees, [101, 102], 2026, 7)
        assert config.MANUAL_EXCLUSIONS == {102: config.MANUAL_EXCLUSION_REASON}
        assert current_settlement_ids([101, 102], setup_employees) == [101]
        assert current_settlement_ids([101, 102], setup_employees, {101: 'sqlite'}) == []

    def test_bad_number_reprompts_and_eof_leaves(self, setup_employees, capsys,
                                                 mock_holidays_jan2026, mock_postponed_empty):
        from core.review import run_review

        with patch('builtins.input', feed('99', 'abc')):  # затем EOF — выход без трейсбека
            run_review(self._table(), setup_employees, [101, 102], 2026, 7)
        assert capsys.readouterr().out.count('Нет сотрудника с номером') == 2

    def test_number_opens_employee_review_and_autofills(self, setup_employees,
                                                        mock_holidays_jan2026, mock_postponed_empty):
        """1 → сотрудник; [2] автозаполнение одиночных; д; [0] назад; Enter — к расчёту."""
        from core.review import run_review

        tt = {f'2026-07-{d:02d}': {101: day(f'2026-07-{d:02d}', '08:00:00', '17:00:00')}
              for d in (6, 7, 8, 9)}
        tt['2026-07-10'] = {101: day('2026-07-10', '08:01:00', '08:01:00')}
        singles = [['2026-07-10', make_dt('2026-07-10', '08:01:00')]]
        state = {'fixed': False}

        def fake_marks(table, emp_id, year, month, employees=None, rules_by_role=None):
            m = table['2026-07-10'][101]
            return (0, 0) if m.go != m.come else (singles, 0)

        with patch.object(analysis, '_get_marks_and_missed', fake_marks), \
                patch.object(analysis, '_save_session', lambda t: state.update(fixed=True)), \
                patch('builtins.input', feed('1', '2', 'д', '')):
            run_review(tt, {101: setup_employees[101]}, [101], 2026, 7)
        assert state['fixed']
        assert tt['2026-07-10'][101].go > tt['2026-07-10'][101].come


# --- Сквозные прогоны main ---------------------------------------------------------

@pytest.fixture
def september_dir(tmp_path):
    root = tmp_path / 'input'
    var = root / 'variable_data_for_app'
    var.mkdir(parents=True)
    (var / 'roles_employee.dat').write_text('[1] Работник цеха\n[4] Ремонтники\n', encoding='utf-8')
    (var / 'id_employee.dat').write_text(
        '101 [1] Иванов Иван\n102 [1] Петров Пётр\n201 [4] Смирнов Максим\n', encoding='utf-8')
    for name in ('settlement_exceptions.dat', 'holidays.dat', 'postponed_working_days.dat'):
        (var / name).write_text('', encoding='utf-8')
    lines = []
    for d in range(1, 31):
        date = datetime(2026, 9, d)
        if date.weekday() >= 5:
            continue
        for emp in (101, 102):
            lines += [f'{emp} {date:%Y-%m-%d} 08:00:00 1', f'{emp} {date:%Y-%m-%d} 17:00:00 1']
    for d in (12, 15, 18, 25):
        lines.append(f'201 2026-09-{d:02d} 09:00:00 1')
    (root / 'marks.dat').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return root


def _argv(root, out, *extra):
    return ['--data-dir', str(root), '--output-dir', str(out), '-f', 'marks.dat',
            '-y', '2026', '-m', '9', '-k', 't', *extra]


def test_main_shows_repairmen_and_excludes_chosen_employee(september_dir, tmp_path,
                                                           monkeypatch, capsys):
    import json

    from openpyxl import load_workbook

    from main import main

    monkeypatch.chdir(tmp_path)
    out = tmp_path / 'out'
    monkeypatch.setattr('builtins.input', feed('x2', ''))
    main(_argv(september_dir, out))
    text = capsys.readouterr().out
    assert 'Максим Смирнов — 12, 15, 18, 25 сентября' in text
    assert 'Ср. приход' in text and '08:00:00' in text
    # в отчёты исключённый не попал: ни HTML, ни Excel, ни ведомость
    assert [p.name for p in out.glob('*.html')] == ['Иванов Иван_101_09_2026.html']
    wb = load_workbook(next(out.glob('*.xlsx')))
    try:
        cells = {c.value for row in wb.active.iter_rows() for c in row if c.value}
    finally:
        wb.close()
    assert 'Иванов Иван' in cells and 'Петров Пётр' not in cells
    package = json.loads(next(out.glob('payroll_2026_09_rev*.json')).read_text(encoding='utf-8'))
    assert package['participants'] == [101]
    assert package['excluded']['102'] == config.MANUAL_EXCLUSION_REASON
    assert config.MANUAL_EXCLUSIONS == {}  # запуск не оставляет состояния в процессе


def test_resume_keeps_manual_exclusion(september_dir, tmp_path, monkeypatch, capsys):
    import main as main_mod

    monkeypatch.chdir(tmp_path)
    out = tmp_path / 'out'
    real_excel = main_mod.build_excel

    def broken_excel(*a, **k):
        raise PermissionError('файл открыт в Excel')

    monkeypatch.setattr(main_mod, 'build_excel', broken_excel)
    monkeypatch.setattr('builtins.input', feed('x2', ''))
    with pytest.raises(SystemExit) as exc:
        main_mod.main(_argv(september_dir, out))
    assert exc.value.code == 1
    assert (out / 'temporary.json').exists()  # сессия оставлена для --resume

    monkeypatch.setattr(main_mod, 'build_excel', real_excel)
    capsys.readouterr()
    main_mod.main(['--resume', '--no-edit', '--data-dir', str(september_dir),
                   '--output-dir', str(out), '-k', 't'])
    assert [p.name for p in out.glob('*.html')] == ['Иванов Иван_101_09_2026.html']
    assert not (out / 'temporary.json').exists()

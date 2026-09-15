"""Отчеты новой модели: один bundle — один результат везде, без своих формул."""

from datetime import timedelta
from decimal import Decimal


def make_dt(date_str, time_str):
    from datetime import datetime
    return datetime.strptime(f'{date_str} {time_str}', '%Y-%m-%d %H:%M:%S')


def make_bundle():
    """Минимальный bundle вручную: без БД, через чистый pay_calc."""
    from core.pay_calc import PayInputs, calculate_pay
    from core.payroll import PayEmployeeResult, PayrollBundle
    from core.roles import get_default_rule

    inputs = PayInputs(
        monthly_base=Decimal('60000.00'), base_day_hours=8,
        full_month_bonus=Decimal('5000.00'), workdays=20,
        shift_norm_hours=Decimal('8'), fact_hours=Decimal('162'),
        workdays_present=20, overtime_eligible=True,
        full_month_eligible=True, seniority_eligible=True,
        overtime_coef=Decimal('1.5'), seniority_rate=Decimal('0.10'),
        milk_amount=Decimal('800.00'),
    )
    result = calculate_pay(inputs)
    bundle = PayrollBundle(
        year=2026, month=10, workdays=20, settings_eff='2026-01-01',
        monthly_base=Decimal('60000.00'), base_day_hours=8,
        full_month_bonus=Decimal('5000.00'), seniority_scale=[],
        rule_versions={1: '2026-01-01'},
        results={101: PayEmployeeResult(
            emp_id=101, name='Петров Иван', rule=get_default_rule(1),
            inputs=inputs, result=result)},
        warnings=['тест-предупреждение'],
    )
    return bundle


def make_legacy_tables():
    """Legacy-таблицы для одного дня (детализация/итоги сохраняются)."""
    from core.day_models import DayMark, DayWork, EmployeeMonth, HolidayGroup, WorkGroup, WageResult
    d = '2026-10-06'
    tt = {d: {101: DayMark(go=make_dt(d, '16:00:00'), come=make_dt(d, '08:00:00'), tag='work')}}
    wt = {d: {101: DayWork(delta=timedelta(0), worked=timedelta(hours=8),
                           overtime_tag='недоработка', day_tag='work')}}
    summary = {101: EmployeeMonth(work=WorkGroup(1, timedelta(0), timedelta(0)),
                                  holiday=HolidayGroup(0, timedelta(0), timedelta(0), timedelta(0)),
                                  vacation_days=0, truancy_days=0)}
    wages = {101: WageResult(Decimal('800.00'), Decimal('40.00'), Decimal('840.00'))}
    return tt, wt, summary, wages


class TestPayHeaderSingleSource:
    def test_header_same_everywhere(self):
        from core.payroll import pay_header_text
        from core.excel_builder import _pay_header_text as excel_h
        from core.html_builder import _pay_header_text as html_h
        bundle = make_bundle()
        assert excel_h(bundle) == pay_header_text(bundle)
        assert html_h(bundle) == pay_header_text(bundle)
        assert '60000.00' in pay_header_text(bundle)
        assert '20' in pay_header_text(bundle)


class TestExcelPayBlock:
    def test_legacy_without_bundle_has_no_pay_block(self, setup_employees, tmp_path, monkeypatch):
        from core.excel_builder import build_excel
        from openpyxl import load_workbook
        monkeypatch.chdir(tmp_path)
        tt, wt, summary, wages = make_legacy_tables()
        f = build_excel(tt, wt, summary, wages, employees=setup_employees)
        wb = load_workbook(tmp_path / f)
        texts = [str(c.value) for row in wb.active.iter_rows() for c in row if c.value]
        assert not any('Новая модель' in t for t in texts)

    def test_bundle_block_shows_single_source_values(self, setup_employees, tmp_path, monkeypatch):
        from core.excel_builder import build_excel
        from openpyxl import load_workbook
        monkeypatch.chdir(tmp_path)
        bundle = make_bundle()
        res = bundle.results[101].result
        tt, wt, summary, wages = make_legacy_tables()
        f = build_excel(tt, wt, summary, wages, employees=setup_employees, bundle=bundle)
        wb = load_workbook(tmp_path / f)
        cells = [c.value for row in wb.active.iter_rows() for c in row if c.value is not None]
        texts = [str(v) for v in cells]
        assert any('Новая модель' in t for t in texts)
        assert any('Петров Иван' in t for t in texts)
        # основания из bundle, без пересчета: сравниваем числа как Decimal,
        # текст (причина/предупреждение) — подстрокой.
        from decimal import Decimal as _D
        nums = {_D(str(v)) for v in cells if isinstance(v, (int, float, _D))}
        for needle in (res.ordinary_pay, res.overtime_bonus,
                       res.full_month_bonus, res.seniority_bonus, res.total):
            assert _D(str(needle)) in nums
        flat = ' '.join(texts)
        for needle in (res.full_month_reason, 'тест-предупреждение'):
            assert needle in flat


class TestHtmlPayBlock:
    def test_legacy_without_bundle_has_no_pay_block(self, setup_employees, tmp_path, monkeypatch):
        from core.html_builder import build_html
        monkeypatch.chdir(tmp_path)
        tt, wt, summary, wages = make_legacy_tables()
        f = build_html(101, tt, wt, summary, wages, employees=setup_employees)
        text = (tmp_path / f).read_text(encoding='utf-8')
        assert 'Новая модель' not in text

    def test_bundle_block_personal_row(self, setup_employees, tmp_path, monkeypatch):
        from core.html_builder import build_html
        from core.money import format_money
        monkeypatch.chdir(tmp_path)
        bundle = make_bundle()
        res = bundle.results[101].result
        tt, wt, summary, wages = make_legacy_tables()
        f = build_html(101, tt, wt, summary, wages, employees=setup_employees, bundle=bundle)
        text = (tmp_path / f).read_text(encoding='utf-8')
        assert 'Новая модель оплаты' in text
        assert 'Петров Иван' in text
        assert format_money(res.total) in text
        assert format_money(res.ordinary_pay) in text
        assert res.full_month_reason in text or 'полный' in text.lower()
        assert 'тест-предупреждение' in text

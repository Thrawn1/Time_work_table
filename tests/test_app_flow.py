"""Регрессии запуска на синтетических справочниках и настоящих файлах отчётов."""

import pytest
from openpyxl import load_workbook

from core import config, file_parser, session
from core.pay_store import assign_role, init_db, seed_defaults, upsert_employee
from main import main


@pytest.fixture
def input_dir(tmp_path):
    root = tmp_path / 'input'
    variable = root / 'variable_data_for_app'
    variable.mkdir(parents=True)
    files = {
        'roles_employee.dat': '[1] Работник\n',
        'id_employee.dat': '101  [1]\tТестов   Тест\n',
        'settlement_exceptions.dat': '',
        'holidays.dat': '',
        'postponed_working_days.dat': '',
    }
    for name, text in files.items():
        (variable / name).write_text(text, encoding='utf-8')
    (root / 'marks.dat').write_text(
        '101 2026-07-06 08:00:00 1 255 1 0\n'
        '101 2026-07-06 16:00:00 1 255 1 0\n', encoding='utf-8')
    return root


def _run_args(input_dir, output_dir):
    return ['--data-dir', str(input_dir), '--output-dir', str(output_dir),
            '-f', 'marks.dat', '-y', '2026', '-m', '7', '--no-edit']


def _summary_values(output_dir):
    wb = load_workbook(next(output_dir.glob('*.xlsx')), data_only=True)
    try:
        return next(row for row in wb.active.iter_rows(values_only=True)
                    if len(row) >= 17 and row[7] == 'Тестов Тест')
    finally:
        wb.close()


def _seed_pay_db(input_dir):
    con = init_db(str(input_dir / 'pay_directory.db'))
    try:
        seed_defaults(con)
        upsert_employee(con, 101, 'Тест', 'Тестов')
        assign_role(con, 101, 1, '2026-01-01')
    finally:
        con.close()


@pytest.mark.parametrize('key', ['t', '0', 'invalid', None])
def test_time_only_with_pay_database(input_dir, tmp_path, monkeypatch, key):
    _seed_pay_db(input_dir)
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr('sys.stdin.isatty', lambda: False)
    output = tmp_path / 'out'
    args = _run_args(input_dir, output)
    if key is not None:
        args.extend(['-k', key])
    # wage_rates.dat отсутствует: для режима учёта чтение ставок не требуется.
    main(args)
    row = _summary_values(output)
    assert row[14:17] == (0, 40, 40)
    assert len(list(output.glob('*.html'))) == 1
    # Новая модель действует на месяц независимо от salary_mode (F02): состав
    # и молоко считаются по SQLite-условиям, версии сохраняются, оклады обнулены.
    assert list(output.glob('payroll_versions_*.json'))


def test_salary_mode_uses_new_model(input_dir, tmp_path, monkeypatch):
    _seed_pay_db(input_dir)
    monkeypatch.chdir(tmp_path)
    output = tmp_path / 'out'
    main(_run_args(input_dir, output) + ['-k', '100'])
    row = _summary_values(output)
    assert row[14] > 0
    assert row[15] == 40
    assert row[16] == pytest.approx(row[14] + 40)
    assert list(output.glob('payroll_versions_*.json'))


def test_payroll_service_uses_explicit_employees(input_dir):
    from core.calculations import calculate_hours_per_day, calculate_hours_per_month
    from core.data_array import build_data_array
    from core.payroll_service import calculate_payroll

    _seed_pay_db(input_dir)
    staff = {101: config.EmployeeData(101, 'Явный', 'Сотрудник', 1, 'Работник')}
    with config.preserve_config():
        config.set_data_dir(str(input_dir))
        config.EMPLOYEES.clear()
        config.SETTLEMENT_EXCEPTIONS.clear()
        lines = file_parser.read_file_data('marks.dat', 2026, 7)
        table = build_data_array(lines, employees=staff)
        work_time = calculate_hours_per_day(table, employees=staff)
        summary, _ = calculate_hours_per_month(work_time)
        outcome = calculate_payroll(table, work_time, summary, 2026, 7,
                                    str(input_dir / 'pay_directory.db'),
                                    salary_mode=True, employees=staff)
        assert outcome.wages[101].salary > 0
        assert outcome.bundle.results[101].name == 'Сотрудник Явный'


def test_bad_pay_database_aborts_without_fallback(input_dir, tmp_path, monkeypatch, capsys):
    """Повреждённый pay_directory.db фатален и не даёт молчаливый откат в legacy (F05)."""
    (input_dir / 'pay_directory.db').write_text('not sqlite', encoding='utf-8')
    (input_dir / 'variable_data_for_app' / 'wage_rates.dat').write_text(
        '101\t [00.800]\n', encoding='utf-8')
    monkeypatch.chdir(tmp_path)
    output = tmp_path / 'out'
    with pytest.raises(SystemExit) as exc:
        main(_run_args(input_dir, output) + ['-k', '100'])
    assert exc.value.code == 1
    assert 'справочник оплаты недоступен' in capsys.readouterr().out
    assert not list(output.glob('*.xlsx'))


def test_calendar_cache_follows_directory_without_manual_reset(input_dir, tmp_path):
    other = tmp_path / 'other'
    variable = other / 'variable_data_for_app'
    variable.mkdir(parents=True)
    (variable / 'holidays.dat').write_text('06.07\n', encoding='utf-8')
    (variable / 'postponed_working_days.dat').write_text('04.07\n', encoding='utf-8')
    with config.preserve_config():
        config.set_data_dir(str(input_dir))
        assert file_parser.definition_of_working_day('2026-07-06')[0] == 'work'
        assert file_parser.definition_of_working_day('2026-07-04')[0] == 'weekend'
        config.set_data_dir(str(other))
        assert file_parser.definition_of_working_day('2026-07-06')[0] == 'holiday'
        assert file_parser.definition_of_working_day('2026-07-04')[0] == 'work'
        config.set_data_dir(str(input_dir))
        assert file_parser.definition_of_working_day('2026-07-06')[0] == 'work'


@pytest.mark.parametrize('missing_file', [False, True])
def test_main_restores_config_on_success_and_failure(
    input_dir, tmp_path, monkeypatch, setup_employees, missing_file,
):
    roles, employees, exceptions = config.ROLES, config.EMPLOYEES, config.SETTLEMENT_EXCEPTIONS
    before = roles.copy(), employees.copy(), exceptions.copy()
    paths = config.DATA_DIR, config._SECRET_KEY, session.SESSION_FILE, session.SESSION_FILE_LEGACY
    monkeypatch.chdir(tmp_path)
    args = _run_args(input_dir, tmp_path / 'out') + ['-k', 't']
    if missing_file:
        args.extend(['-f', 'missing.dat'])
        with pytest.raises(SystemExit) as exc:
            main(args)
        assert exc.value.code == 1
    else:
        main(args)
    assert config.ROLES is roles and config.EMPLOYEES is employees
    assert config.SETTLEMENT_EXCEPTIONS is exceptions
    assert (roles, employees, exceptions) == before
    assert (config.DATA_DIR, config._SECRET_KEY, session.SESSION_FILE,
            session.SESSION_FILE_LEGACY) == paths


def test_failed_config_load_is_atomic(input_dir, setup_employees):
    with config.preserve_config():
        config.set_data_dir(str(input_dir))
        (input_dir / 'variable_data_for_app' / 'settlement_exceptions.dat').write_text(
            'invalid-id\n', encoding='utf-8')
        before = config.ROLES.copy(), config.EMPLOYEES.copy(), config.SETTLEMENT_EXCEPTIONS.copy()
        with pytest.raises(config.ConfigError):
            config.load_config()
        assert (config.ROLES, config.EMPLOYEES, config.SETTLEMENT_EXCEPTIONS) == before


def test_missing_resume_does_not_start_new_import(input_dir, tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    output = tmp_path / 'out'
    with pytest.raises(SystemExit) as exc:
        main(_run_args(input_dir, output) + ['--resume', '-k', 't'])
    assert exc.value.code == 2
    assert '--resume не найдена' in capsys.readouterr().err
    assert not list(output.glob('*.xlsx'))
